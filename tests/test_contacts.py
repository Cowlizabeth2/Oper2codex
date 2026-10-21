from openpyxl import load_workbook

from operational_center.contacts import create_directory, ContactsService


def test_formatted_directory(tmp_path):
    path = tmp_path / "contacts.xlsx"
    create_directory(path)
    wb = load_workbook(path)
    try:
        ws = wb.active
        assert ws.freeze_panes == "A2"
        assert ws.auto_filter.ref == "A1:H1001"
        assert ws["B2"].number_format == "@"
        assert ws["A1"].font.bold
        assert ws.column_dimensions["B"].width >= 20
    finally:
        wb.close()


def test_exact_tags_optional_fields_actions_update_and_cache(store, tmp_path):
    path = tmp_path / "contacts.xlsx"
    create_directory(path)
    wb = load_workbook(path)
    ws = wb.active
    ws.append(["Иванов", "+70000000001", None, " ДТП , Пожар ", "Позвонить, Написать, Отправить файл", None])
    ws.append(["Петров", "002", None, "ДТП 2", None, None])
    ws.append(["Без телефона", None, None, "ДТП"])
    wb.save(path); wb.close()
    service = ContactsService(store)
    contacts, warning = service.load(path, "dtp", {"дтп": "dtp"})
    assert not warning and len(contacts) == 1
    assert contacts[0]["name"] == "Иванов" and contacts[0]["role"] == ""
    assert len(contacts[0]["actions"]) == 3
    assert len(service.load(path, "ПОЖАР")[0]) == 1
    wb = load_workbook(path)
    # append() starts after formatted blank rows, which are intentionally ignored.
    ws = wb.active
    ws.append(["Сидоров", "003", None, "ДТП"])
    wb.save(path); wb.close()
    assert len(service.load(path, "dtp", {"ДТП": "dtp"}, True)[0]) == 2
    path.unlink()
    stale, warning = service.load(path, "dtp", {"ДТП": "dtp"})
    assert len(stale) == 2 and "неактуальный кэш" in warning


def test_existing_excel_never_overwritten(tmp_path):
    import pytest
    path = tmp_path / "existing.xlsx"
    path.write_bytes(b"keep")
    with pytest.raises(FileExistsError):
        create_directory(path)
    assert path.read_bytes() == b"keep"
