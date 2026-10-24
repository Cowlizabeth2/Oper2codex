"""Synthetic review data, created only in a caller-provided temporary directory."""
from pathlib import Path
from openpyxl import load_workbook
from operational_center.contacts import create_directory
from operational_center.domain import new_block, field


def prepare(store, root):
    root = Path(root)
    contacts = Path(create_directory(root / 'Демонстрационные_контакты.xlsx'))
    wb = load_workbook(contacts)
    rows = [
        ['Иванов Иван Иванович', '+7 (000) 000-00-01', 'Дежурный диспетчер', 'ДТП, Пожар', 'Позвонить, Написать', 'Передать адрес и характер происшествия'],
        ['Петрова Анна Сергеевна', '+7 (000) 000-00-02', 'Оперативный дежурный', 'ДТП', 'Позвонить, Отправить файл', 'Направить заполненную карточку'],
        ['Смирнов Алексей Павлович', '+7 (000) 000-00-03', '', 'ДТП', 'Написать', ''],
    ]
    for row_index, row in enumerate(rows, 2):
        for column, value in enumerate(row, 1):
            wb.active.cell(row_index, column, value)
    wb.save(contacts); wb.close()
    files=[]
    for name, text in [('Порядок оповещения.txt', 'Демонстрационные данные для проверки интерфейса.'), ('Памятка сотруднику.txt', 'Проверить адрес. Заполнить карточку. Уведомить дежурного.')]:
        p=root/name;p.write_text(text, encoding='utf-8');files.append(str(p))
    time=field('Время', default='14:35')
    place=field('Место', default='ул. Примерная, 12')
    carrier=field('Перевозчик', default='ТМП20', phrases=['ТМП20','Стартранс','Автолайн'])
    detail=field('Происшествие', default='Столкновение двух автомобилей.', multiline=True)
    injured=field('Пострадавшие', default='Нет', phrases=['Нет','Уточняется'], required=True)
    text=new_block('text')
    text.update(title='Оперативное сообщение', x=260,y=40,width=560,height=420)
    text['config']['segments']=[{'kind':'text','text':'Время: '},time,{'kind':'text','text':'  Место: '},place,{'kind':'text','text':'\nПеревозчик: '},carrier,{'kind':'text','text':'\nПроисшествие: '},detail,{'kind':'text','text':'\nПострадавшие: '},injured]
    people=new_block('contacts');people.update(x=844,y=40,width=400,height=300,title='Контакты оповещения');people['config'].update(path=str(contacts),tag_map={'ДТП':'dtp','Пожар':'fire'})
    docs=new_block('documents');docs.update(x=260,y=484,width=560,height=288,title='Рабочий комплект');docs['config']['sources']=[str(contacts),files[0]]
    refs=new_block('reference');refs.update(x=844,y=364,width=400,height=408,title='Сопутствующие файлы');refs['config']['sources']=files+[str(contacts)]
    _,config=store.config('draft')
    master=config['masters'][0]
    master.update(blocks=[text,people,docs,refs],links=[b['id'] for b in [text,people,docs,refs]])
    store.save_draft(config);store.publish()
    return master, text
