"""Excel directories with exact, normalized tags and a persistent cache."""
import json
import logging
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

from .domain import normalize_tag
from .storage import encode

HEADERS = ["ФИО", "Телефон", "Должность", "Теги", "Действия", "Комментарий", "Приоритет", "Срок оповещения"]


def create_directory(path):
    path = Path(path)
    if path.exists():
        raise FileExistsError("Файл уже существует; выберите другое имя")
    wb = Workbook()
    ws = wb.active
    ws.title = "Контакты"
    ws.append(HEADERS)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:H1001"
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="263A52")
        cell.alignment = Alignment(wrap_text=True)
    for column, width in zip("ABCDEFGH", [32, 24, 28, 24, 36, 48, 16, 24]):
        ws.column_dimensions[column].width = width
    for row in range(2, 1002):
        ws.cell(row, 2).number_format = "@"
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        with path.open("xb") as stream:
            wb.save(stream)
    finally:
        wb.close()
    return str(path)


class ContactsService:
    def __init__(self, store):
        self.store = store

    def load(self, path, tag, mapping=None, refresh=False):
        if not path:
            return [], "Справочник не подключён"
        absolute = str(Path(path).expanduser().resolve())
        cache = self.store.db.execute("SELECT stamp,data FROM contacts_cache WHERE path=?", (absolute,)).fetchone()
        warning = ""
        try:
            stat = Path(absolute).stat()
            stamp = f"{stat.st_mtime_ns}:{stat.st_size}"
            if cache and cache["stamp"] == stamp and not refresh:
                rows = json.loads(cache["data"])
            else:
                wb = load_workbook(absolute, read_only=True, data_only=True)
                try:
                    iterator = iter(wb.active.iter_rows(values_only=True))
                    columns = [str(c or "").strip() for c in next(iterator)]
                    if not {"ФИО", "Телефон"} <= set(columns):
                        raise ValueError("В Excel обязательны колонки ФИО и Телефон")
                    rows = []
                    for values in iterator:
                        row = {key: str(value if value is not None else "").strip() for key, value in zip(columns, values)}
                        if not row.get("ФИО") or not row.get("Телефон"):
                            continue
                        rows.append(row)
                finally:
                    wb.close()
                with self.store.db:
                    self.store.db.execute("INSERT OR REPLACE INTO contacts_cache VALUES(?,?,?)", (absolute, stamp, encode(rows)))
        except Exception as exc:
            logging.getLogger(__name__).exception("Ошибка чтения справочника %s", absolute)
            rows = json.loads(cache["data"]) if cache else []
            warning = f"Справочник недоступен: {exc}. " + ("Показан неактуальный кэш." if cache else "Нет данных.")
        aliases = {normalize_tag(k): normalize_tag(v) for k, v in (mapping or {}).items()}
        selected = []
        for row in rows:
            tags = {aliases.get(normalize_tag(t), normalize_tag(t)) for t in row.get("Теги", "").split(",") if t.strip()}
            if normalize_tag(tag) in tags:
                selected.append({"name": row["ФИО"], "phone": row["Телефон"], "role": row.get("Должность", ""),
                                 "comment": row.get("Комментарий", ""), "actions": [a.strip() for a in row.get("Действия", "").split(",") if a.strip()],
                                 "priority": row.get("Приоритет", ""), "deadline": row.get("Срок оповещения", "")})
        return sorted(selected, key=lambda r: r["name"].casefold()), warning
