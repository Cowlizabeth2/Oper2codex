"""Explicit entrypoint, startup diagnostics and packaged QML smoke-test."""
from __future__ import annotations

import argparse
import faulthandler
import json
import logging
import os
import sys
import tempfile
from pathlib import Path

from PySide6.QtCore import QTimer, QUrl, qInstallMessageHandler, QtMsgType, QMetaObject, Qt, QObject
from PySide6.QtQml import QQmlApplicationEngine
from PySide6.QtQuick import QQuickWindow
from PySide6.QtWidgets import QApplication, QMessageBox

from . import __version__
from .controller import Controller
from .storage import Store


def default_data_dir():
    base = Path(os.environ.get("LOCALAPPDATA", Path.home() / ".local" / "share"))
    return base / "OperationalCenter"


def main(argv=None):
    parser = argparse.ArgumentParser(description="Оперативный центр")
    parser.add_argument("--data-dir", type=Path)
    parser.add_argument("--smoke-test", action="store_true")
    parser.add_argument('--upgrade-test',action='store_true')
    parser.add_argument('--upgrade-manifest',type=Path)
    parser.add_argument('--upgrade-phase',choices=['apply','verify'],default='apply')
    parser.add_argument('--upgrade-output',type=Path)
    parser.add_argument("--smoke-report", type=Path)
    parser.add_argument("--screenshot", type=Path)
    parser.add_argument("--screenshot-page", default="home", choices=["home", "builder", "events"])
    args = parser.parse_args(argv)
    upgrade_manifest=None
    if args.upgrade_test:
        if args.smoke_test or not args.data_dir or not args.upgrade_manifest or not args.upgrade_output:
            parser.error('Upgrade test requires an explicit synthetic data dir, manifest and output')
        from .upgrade_check import guard
        upgrade_manifest=guard(args.data_dir,args.upgrade_manifest,args.upgrade_phase)
    checking=args.smoke_test or args.upgrade_test
    # Smoke tests NEVER open or migrate the user's database.
    temporary = tempfile.TemporaryDirectory(prefix="oper-smoke-") if args.smoke_test else None
    data_dir = Path(temporary.name) if temporary else (args.data_dir or default_data_dir())
    data_dir.mkdir(parents=True, exist_ok=True)
    startup = data_dir / "startup.log"
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(name)s: %(message)s",
                        handlers=[logging.FileHandler(startup, encoding="utf-8"), logging.FileHandler(data_dir / "application.log", encoding="utf-8"), logging.StreamHandler()], force=True)
    crash_stream = (data_dir / "crash.log").open("a", encoding="utf-8")
    faulthandler.enable(crash_stream)
    qml_errors = []
    previous_hook = sys.excepthook

    def unhandled(exc_type, value, traceback):
        logging.critical("Необработанное исключение", exc_info=(exc_type, value, traceback))
        if checking:
            app.exit(4)
        else:
            QMessageBox.critical(None, "Ошибка приложения", f"{value}\n\nПодробности: {startup}")

    sys.excepthook = unhandled

    def qt_message(kind, context, message):
        logging.getLogger("Qt").log(logging.WARNING if kind in {QtMsgType.QtWarningMsg, QtMsgType.QtCriticalMsg, QtMsgType.QtFatalMsg} else logging.INFO, message)
        if kind in {QtMsgType.QtWarningMsg, QtMsgType.QtCriticalMsg, QtMsgType.QtFatalMsg}:
            qml_errors.append(message)

    qInstallMessageHandler(qt_message)
    app = QApplication.instance() or QApplication(sys.argv[:1])
    app.setApplicationName("Оперативный центр")
    app.setOrganizationName("OperationalCenter")
    app.setApplicationVersion(__version__)
    store = None
    result = {"ok": False, "version": __version__, "checks": [], "qml_errors": qml_errors}
    try:
        build_file = Path(__file__).parent / "build.json"
        commit = json.loads(build_file.read_text(encoding="utf-8")).get("commit", "development") if build_file.exists() else "development"
        result['commit'] = commit
        logging.info("Оперативный центр %s, build %s; data=%s", __version__, commit, data_dir)
        store = Store(data_dir / "operational-center.sqlite")
        backend = Controller(store, data_dir, commit)
        engine = QQmlApplicationEngine()
        engine.rootContext().setContextProperty("backend", backend)
        engine.load(QUrl.fromLocalFile(str(Path(__file__).parent / "qml" / "Main.qml")))
        if not engine.rootObjects():
            raise RuntimeError("Не удалось загрузить интерфейс QML")
        root = engine.rootObjects()[0]
        smoke_steps = []
        if args.upgrade_test:
            def upgrade_run():
                try:
                    from .upgrade_check import run
                    result['upgrade']=run(backend,root,upgrade_manifest,args.upgrade_phase,data_dir,args.upgrade_output)
                    result['ok']=not qml_errors
                    app.exit(0 if result['ok'] else 2)
                except Exception as exc:
                    logging.exception('Packaged upgrade test failed')
                    result['error']=str(exc);app.exit(2)
            QTimer.singleShot(200,upgrade_run)
            QTimer.singleShot(90000,lambda:app.exit(3))
        if args.smoke_test:
            # Exercise all QML pages and each functional block using temporary data.
            from .domain import new_block, field
            from .contacts import create_directory
            master = store.config("draft")[1]["masters"][0]
            source = data_dir / "source.txt"
            source.write_text("Безопасная проверка документов", encoding="utf-8")
            contacts = create_directory(data_dir / "contacts.xlsx")
            from openpyxl import load_workbook
            workbook = load_workbook(contacts)
            for column, value in enumerate(["Проверочный контакт", "+70000000000", "", "dtp", "Позвонить"], 1):
                workbook.active.cell(2, column, value)
            workbook.save(contacts)
            workbook.close()
            blocks = [new_block(kind) for kind in ("text", "contacts", "documents", "reference", "static_text")]
            for index, block in enumerate(blocks):
                block.update(x=260 + index % 2 * 500, y=60 + index // 2 * 390)
            blocks[0]["config"]["segments"] = [{"kind": "text", "text": "Проверка: "}, field("Сообщение", default="тест")]
            blocks[1]["config"]["path"] = contacts
            for block in blocks[2:4]:
                block["config"]["sources"] = [str(source)]
            blocks[4]["config"].update(text="**Памятка**\nСтрока 2", format="markdown")
            _, config = store.config("draft")
            config["masters"][0]["blocks"] = blocks
            config["masters"][0]["links"] = [b["id"] for b in blocks]
            from .catalog import new_catalog
            catalog = new_catalog()
            catalog["members"] = [master["id"]]
            config["catalogs"] = [catalog]
            store.save_draft(config)
            store.publish()
            if not backend.command("launch_catalog", json.dumps({"id": master["id"], "catalog": catalog["id"]})):
                raise RuntimeError("Smoke: создание события не выполнено")
            runtime = backend.state["event"]
            if len(runtime["snapshot"]["blocks"]) != 5:
                raise RuntimeError("Smoke: не все блоки активированы")
            by_type = {b["type"]: b for b in runtime["snapshot"]["blocks"]}
            if by_type["text"]["preview"] != "Проверка: тест" or len(by_type["contacts"]["contacts"]) != 1:
                raise RuntimeError("Smoke: шаблон или импорт Excel не работает")
            copied = by_type["documents"]["files"][0]
            if copied["status"] != "ready" or Path(copied["path"]).read_text(encoding="utf-8") != source.read_text(encoding="utf-8"):
                raise RuntimeError("Smoke: рабочая копия не создана")
            if len(by_type["reference"]["files"]) != 1:
                raise RuntimeError("Smoke: справочный файл не найден")
            result["service_checks"] = ["event_snapshot", "template_render", "excel_import", "document_copy", "reference_scan"]
            if by_type["static_text"]["static"]["plain"] != "Памятка\nСтрока 2" or not backend.searchCatalog(catalog["id"], master["tag"], "", "all"):
                raise RuntimeError("Smoke: свободный текст или каталог не работает")
            from .scenario_import import sample
            sample_path = sample(data_dir / "scenario-import.xlsx")
            if not backend.command("preview_scenarios", json.dumps({"path":sample_path})) or backend.state["importPreview"]["valid"] != 1:
                raise RuntimeError("Smoke: предварительная проверка сценариев не работает")
            result["service_checks"].extend(["static_text", "catalog_search_launch", "scenario_import_preview"])
            if not backend.command('inspect_templates','{}') or not backend.state['templateDiagnostics']['candidates']:
                raise RuntimeError('Smoke: диагностика сохранённых шаблонов недоступна')
            result['service_checks'].append('template_migration_inspection')
            backend.command("select_block", json.dumps({"id": blocks[0]["id"]}))
            smoke_steps = [("events", "text"), ("builder", "text"), ("builder", "contacts"), ("builder", "documents"), ("builder", "reference"), ("builder", "static_text"), ("history", "reference"), ("settings", "reference"), ("logs", "reference"), ("home", "reference"), ("catalog", "search"), ("catalog", "manager"), ("catalog", "import")]
            smoke_steps.append(('migration','review'))
            QTimer.singleShot(15000, lambda: app.exit(3))

        def smoke_next():
            if smoke_steps:
                page, kind = smoke_steps.pop(0)
                for name in ("catalogSearchDialog", "catalogManagerDialog", "scenarioImportDialog", "templateMigrationDialog"):
                    popup = root.findChild(QObject, name)
                    if popup and popup.property("visible"):
                        QMetaObject.invokeMethod(popup, "close", Qt.DirectConnection)
                if page == "catalog":
                    QMetaObject.invokeMethod(root, {"search":"openAllScenarios", "manager":"manageCatalogs", "import":"importScenarios"}[kind], Qt.DirectConnection)
                if page == 'migration':
                    QMetaObject.invokeMethod(root,'inspectTemplates',Qt.DirectConnection)
                if page == "builder":
                    block = next(b for b in blocks if b["type"] == kind)
                    backend.command("select_block", json.dumps({"id": block["id"]}))
                if page not in {'catalog','migration'}:
                    root.setProperty("page", page)
                result["checks"].append(page + ":" + kind)
                QTimer.singleShot(150, smoke_next)
            else:
                if args.screenshot:
                    root.setProperty("page", args.screenshot_page)
                    app.processEvents()
                    root.grabWindow().save(str(args.screenshot))
                result["ok"] = not qml_errors
                app.exit(0 if result["ok"] else 2)

        if args.smoke_test:
            QTimer.singleShot(250, smoke_next)
        code = app.exec()
        # Destroy the engine while context properties still exist.
        del root
        del engine
        if checking:
            result["ok"] = code == 0 and not qml_errors
            logging.info("SMOKE_RESULT %s", json.dumps(result, ensure_ascii=False))
        return code
    except Exception as exc:
        logging.exception("Ошибка запуска")
        result["error"] = str(exc)
        if not checking:
            QMessageBox.critical(None, "Ошибка запуска — Оперативный центр", f"{exc}\n\nДиагностика: {startup}")
        return 1
    finally:
        if args.smoke_report:
            args.smoke_report.parent.mkdir(parents=True, exist_ok=True)
            if not result["ok"]:
                result["startup_log"] = startup.read_text(encoding="utf-8")
            args.smoke_report.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
        if store:
            store.close()
        qInstallMessageHandler(None)
        sys.excepthook = previous_hook
        faulthandler.disable()
        crash_stream.close()
        # Windows refuses to remove open diagnostic files in a smoke workspace.
        logging.shutdown()
        if temporary:
            temporary.cleanup()
