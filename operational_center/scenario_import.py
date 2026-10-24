"""Validate Excel into an immutable proposal. Applying only appends new scenarios."""
from __future__ import annotations

import copy
import hashlib
import json
import re
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .catalog import normalize
from .domain import duplicate_master, new_block, validate_config

HEADERS = {
    "name": ["название", "название станции или происшествия", "станция"],
    "code": ["системный код", "уникальный системный код"],
    "tags": ["поисковые теги", "теги"],
    "aliases": ["альтернативные названия", "алиасы"],
    "category": ["категория", "категория или линия", "линия"],
    "documents": ["путь к папке документов", "папка документов"],
    "template": ["шаблон мастер-сценария", "шаблон", "сценарий-шаблон", "шаблон мастер-сценария, который нужно использовать"],
    "km_code": ["код км", "км"],
}


def fingerprint(config):
    return hashlib.sha256(json.dumps(config, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def split_values(value):
    return list(dict.fromkeys(v.strip() for v in re.split(r"[;|\n]", value) if v.strip()))


def sample(path):
    path = Path(path)
    book = Workbook()
    sheet = book.active
    sheet.title = "Сценарии"
    sheet.append([values[0] for values in HEADERS.values()])
    sheet.append(["Станция Пример", "station_001", "метро;пересадка", "Примерная", "Линия 1", "", "dtp", "КМ 1"])
    try:
        with path.open("xb") as stream:
            book.save(stream)
    finally:
        book.close()
    return str(path)


def prepare(path, config):
    path = Path(path)
    if path.suffix.lower() != ".xlsx" or not path.is_file() or path.stat().st_size > 10 * 1024 * 1024:
        raise ValueError("Выберите XLSX размером до 10 МБ")
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = book.active
        if sheet.max_row > 5001 or sheet.max_column > 100:
            raise ValueError("Не более 5000 строк и 100 столбцов за импорт")
        rows = sheet.iter_rows()
        first = next(rows, [])
        mapping = {}
        for index, cell in enumerate(first):
            for key, aliases in HEADERS.items():
                if normalize(cell.value or "") in [normalize(key), *map(normalize, aliases)]:
                    if key in mapping:
                        raise ValueError("Повторяющийся столбец: " + aliases[0])
                    mapping[key] = index
        if not {"name", "code", "template"} <= mapping.keys():
            raise ValueError("Обязательные столбцы: название, системный код, шаблон мастер-сценария")
        raw = []
        for number, cells in enumerate(rows, 2):
            values = {key: str(cells[index].value if cells[index].value is not None else "").strip() for key, index in mapping.items()}
            if not any(values.values()):
                continue
            formula = any(cells[index].data_type == "f" for index in mapping.values())
            raw.append((number, values, formula))
        counts = Counter(normalize(v.get("code", "")) for _, v, _ in raw)
        existing = {normalize(m["tag"]) for m in config["masters"]}
        result, proposed = [], []
        for number, values, formula in raw:
            errors = []
            name, code, template_key = (values.get(key, "") for key in ("name", "code", "template"))
            if not name or not code or not template_key:
                errors.append("Не заполнены название, код или шаблон")
            if formula:
                errors.append("Формулы не поддерживаются: вставьте значения")
            if normalize(code) in existing:
                errors.append("Системный код уже существует")
            if counts[normalize(code)] > 1:
                errors.append("Дубликат системного кода в таблице")
            templates = [m for m in config["masters"] if template_key == m["id"] or normalize(template_key) == normalize(m["tag"])]
            if len(templates) != 1:
                errors.append("Не найден однозначный шаблон: укажите его системный код или UUID")
            folder = values.get("documents", "")
            if folder and not Path(folder).expanduser().is_dir():
                errors.append("Папка документов не найдена или недоступна")
            master = None
            if not errors:
                master = duplicate_master(templates[0])
                master.update(name=name, tag=code, tags=split_values(values.get("tags", "")),
                              aliases=split_values(values.get("aliases", "")), category=values.get("category", ""),
                              km_code=values.get("km_code", ""))
                if folder:
                    folder = str(Path(folder).expanduser().resolve())
                    documents = [b for b in master["blocks"] if b["type"] == "documents"]
                    if not documents:
                        block = new_block("documents")
                        block.update(x=232, y=40 + len(master["blocks"]) * 312)
                        master["blocks"].append(block)
                        master["links"].append(block["id"])
                        documents = [block]
                    for block in documents:
                        block["config"]["sources"] = [folder]
                # Validate every inherited source before proposing a scenario.
                for block in master["blocks"]:
                    cfg = block["config"]
                    paths = cfg.get("sources", []) if block["type"] in {"documents", "reference"} else [cfg["path"]] if block["type"] == "contacts" and cfg.get("path") else []
                    for source in paths:
                        if not Path(source).expanduser().exists():
                            errors.append("Источник шаблона не найден: " + source)
                try:
                    validate_config({"format": 1, "masters": [master]})
                except (ValueError, KeyError, TypeError) as exc:
                    errors.append(str(exc))
            if not errors:
                proposed.append(master)
            result.append({"row": number, "name": name, "code": code, "category": values.get("category", ""),
                           "km_code": values.get("km_code", ""), "aliases": values.get("aliases", ""), "tags": values.get("tags", ""),
                           "documents": folder, "template": template_key, "errors": errors, "valid": not errors})
        return {"fingerprint": fingerprint(config), "rows": result, "masters": proposed,
                "valid": len(proposed), "invalid": len(result) - len(proposed), "path": str(path)}
    finally:
        book.close()


def apply(config, proposal, catalog_id=""):
    if fingerprint(config) != proposal["fingerprint"]:
        raise ValueError("Черновик изменился. Повторите предварительную проверку Excel")
    if not proposal["masters"]:
        raise ValueError("Нет корректных строк для импорта")
    if any(row["valid"] and row["documents"] and not Path(row["documents"]).is_dir() for row in proposal["rows"]):
        raise ValueError("Папка документов недоступна. Повторите проверку Excel")
    # Sources can disappear between preview and confirmation; fail before writing.
    for master in proposal["masters"]:
        for block in master["blocks"]:
            cfg = block["config"]
            paths = cfg.get("sources", []) if block["type"] in {"documents", "reference"} else [cfg["path"]] if block["type"] == "contacts" and cfg.get("path") else []
            if any(not Path(p).expanduser().exists() for p in paths):
                raise ValueError("Источник недоступен. Повторите проверку Excel")
    config["masters"].extend(copy.deepcopy(proposal["masters"]))
    if catalog_id:
        catalog = next(c for c in config.get("catalogs", []) if c["id"] == catalog_id)
        catalog["members"].extend(m["id"] for m in proposal["masters"])
    validate_config(config)
